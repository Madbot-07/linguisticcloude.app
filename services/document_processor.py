"""
Document Processor
------------------
Extracts plain text from uploaded documents (TXT, MD, CSV, PDF, DOCX)
and rebuilds translated output files for download.
"""

import io

from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {"txt", "md", "csv", "pdf", "docx"}


def allowed_file(filename: str) -> bool:
    """Check whether the uploaded file has a supported extension."""
    return "." in filename and get_extension(filename) in ALLOWED_EXTENSIONS


def get_extension(filename: str) -> str:
    """Return the lowercase extension of a filename (without the dot)."""
    return filename.rsplit(".", 1)[-1].lower()


def extract_text(file_stream, extension: str) -> str:
    """Extract plain text from an uploaded file stream based on its type."""
    if extension in ("txt", "md", "csv"):
        return _extract_plain(file_stream)
    if extension == "pdf":
        return _extract_pdf(file_stream)
    if extension == "docx":
        return _extract_docx(file_stream)
    raise ValueError(f"Unsupported file type: .{extension}")


def _extract_plain(file_stream) -> str:
    raw = file_stream.read()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, AttributeError):
            continue
    raise ValueError("Could not decode text file. Please use UTF-8 encoding.")


def _extract_pdf(file_stream) -> str:
    reader = PdfReader(file_stream)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password-protected and cannot be read.")
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def _extract_docx(file_stream) -> str:
    document = Document(file_stream)
    blocks = []
    for paragraph in document.paragraphs:
        blocks.append(paragraph.text)
    # Also pull text out of tables so nothing is lost
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def build_txt(text: str) -> io.BytesIO:
    """Package translated text as a UTF-8 .txt file."""
    buffer = io.BytesIO(text.encode("utf-8"))
    buffer.seek(0)
    return buffer


def build_docx(text: str, title: str = "") -> io.BytesIO:
    """Package translated text as a .docx file, one paragraph per line."""
    document = Document()
    if title:
        document.add_heading(title, level=1)
    for line in text.split("\n"):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
